"""Rules from the Looker Studio upload docs, plus the workbook round trip."""
import csv
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.deck.dashboard_workbook import SHEETS, TEXT_COLUMNS, build

DIR = Path("outputs/dashboard")
CSVS = sorted(DIR.glob("dashboard_*.csv"))
HEADER_OK = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,127}$")


def read(path):
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.reader(fh))


def test_expected_csv_inventory():
    assert {p.stem for p in CSVS} == set(SHEETS.values())


@pytest.mark.parametrize("path", CSVS, ids=lambda p: p.stem)
def test_csv_rules(path):
    raw = path.read_bytes()
    assert not raw.startswith(b"\xef\xbb\xbf")  # no BOM
    raw.decode("utf-8")
    rows = read(path)
    header = rows[0]
    assert len(rows) > 1, "empty file"
    assert len(set(header)) == len(header), "duplicate headers"
    assert all(HEADER_OK.match(h) for h in header), header
    assert len({len(r) for r in rows}) == 1, "ragged rows"
    for r in rows:
        for cell in r:
            assert "\n" not in cell and "\r" not in cell  # no line breaks inside a field
            assert '"' not in cell  # docs: double quotes inside fields are unsupported
    assert all(any(c != "" for c in col) for col in zip(*rows[1:], strict=True)), "all-empty column"


def test_schemas_are_distinct_so_one_dataset_cannot_hold_them_all():
    headers = {p.stem: tuple(read(p)[0]) for p in CSVS}
    assert len(set(headers.values())) == len(headers)  # every file needs its own dataset or worksheet


def test_workbook_round_trip_preserves_values_blanks_and_unknown(tmp_path):
    out = build(tmp_path / "w.xlsx")
    wb = load_workbook(out, data_only=True)
    assert wb.sheetnames == list(SHEETS)
    assert all(len(n) <= 31 for n in wb.sheetnames)
    for sheet, stem in SHEETS.items():
        rows = read(DIR / f"{stem}.csv")
        ws = wb[sheet]
        got = [[c.value for c in r] for r in ws.iter_rows()]
        assert len(got) == len(rows), sheet
        assert got[0] == rows[0], sheet
        for ri, (want_row, got_row) in enumerate(zip(rows[1:], got[1:], strict=True), start=1):
            for ci, (w, g) in enumerate(zip(want_row, got_row, strict=True)):
                col = rows[0][ci]
                if w == "":
                    assert g is None, (sheet, ri, col, g)  # blank stays blank, never 0 or 'nan'
                elif col in TEXT_COLUMNS.get(sheet, ()):
                    assert isinstance(g, str) and g == w, (sheet, ri, col, g, w)
                else:
                    try:
                        assert float(g) == pytest.approx(float(w)), (sheet, ri, col, g, w)
                    except (TypeError, ValueError):
                        assert str(g) == w, (sheet, ri, col, g, w)


def test_unknown_is_text_not_zero_in_mixed_columns(tmp_path):
    wb = load_workbook(build(tmp_path / "w.xlsx"), data_only=True)
    ws = wb["competitors_matrix"]
    cells = [c.value for r in ws.iter_rows(min_row=2, max_col=6, min_col=2) for c in r]
    assert "unknown" in cells and all(isinstance(v, str) for v in cells)  # scores are text, never numeric
    inr = [r[4].value for r in wb["pricing"].iter_rows(min_row=2)]
    assert all(isinstance(v, str) for v in inr) and "unknown" in inr


def test_month_label_and_suppressed_months_blank():
    rows = read(DIR / "dashboard_reviews_monthly.csv")
    hdr = rows[0]
    assert "month_label" in hdr
    by = {r[hdr.index("month_label")]: r for r in rows[1:]}
    assert list(by) == sorted(by) and all(re.fullmatch(r"\d{4}-\d{2}", m) for m in by)
    july = by["2026-07"]
    assert july[hdr.index("status")] == "suppressed (n<30)"
    assert all(july[hdr.index(c)] == "" for c in hdr if c.endswith("_pct"))  # blank, never 0
    aug = by["2026-08"]
    assert all(aug[hdr.index(c)] != "" for c in hdr if c.endswith("_pct"))


def test_pricing_chart_has_no_missing_prices_and_themes_have_denominators():
    chart = read(DIR / "dashboard_pricing_chart.csv")
    assert all(r[chart[0].index("inr_per_month")] not in ("", "unknown") for r in chart[1:])
    themes = read(DIR / "dashboard_reviews_themes.csv")
    h = themes[0]
    assert {"reviews_all", "reviews_low_rated", "reviews_selected", "coverage_note"} <= set(h)
    assert len(themes) - 1 == 17 and themes[1][h.index("reviews_all")] == "321"
    for r in themes[1:]:
        if r[h.index("share_all_label")].startswith("suppressed"):
            assert r[h.index("share_all_pct")] == ""


def test_summary_and_headline_values_match_ledger_and_sizing():
    s = read(DIR / "dashboard_summary.csv")
    h = s[0]
    assert [float(r[h.index("share_pct")]) for r in s[1:]] == [14.0, 2.0]
    assert float(s[1][h.index("residual_min")]) == 0.223 and float(s[1][h.index("residual_max")]) == 1.224
    e = read(DIR / "dashboard_experiments_headline.csv")
    eh = e[0]
    assert [r[eh.index("total_days_default")] for r in e[1:]] == ["11.2", "25.5", "9.5"]
    assert all(r[eh.index("decision_rule")] for r in e[1:])


def test_no_unplanned_mixed_type_columns():
    """Looker Studio wants one type per column; mixed columns must be declared text."""
    for sheet, stem in SHEETS.items():
        rows = read(DIR / f"{stem}.csv")
        for ci, col in enumerate(rows[0]):
            vals = [r[ci] for r in rows[1:] if r[ci] != ""]
            kinds = set()
            for v in vals:
                try:
                    float(v)
                    kinds.add("num")
                except ValueError:
                    kinds.add("text")
            if kinds == {"num", "text"}:
                assert col in TEXT_COLUMNS.get(sheet, ()), f"{sheet}.{col} mixes numbers and text"
