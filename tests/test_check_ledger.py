from pptx import Presentation
from pptx.util import Inches

from src.common.check_ledger import find_orphans
from src.common.ledger import add_claim


def make_deck(path, text):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = text
    p.save(path)


def test_orphans_detected(tmp_path):
    led = tmp_path / "l.csv"
    add_claim("R-1", "s2", "install share", 14, "%", "unknown", caveat="m", path=led)
    deck = tmp_path / "d.pptx"
    make_deck(deck, "India is 14% of installs and 2% of revenue. Data as of 2026-10-18")
    assert find_orphans(deck, led) == ["2"]


def test_clean_deck_passes(tmp_path):
    led = tmp_path / "l.csv"
    add_claim("R-1", "s2", "install share", 14, "%", "unknown", caveat="m", path=led)
    deck = tmp_path / "d.pptx"
    make_deck(deck, "About 14% of installs")
    assert find_orphans(deck, led) == []


def test_ids_years_and_dates_are_ignored(tmp_path):
    led = tmp_path / "l.csv"
    add_claim("R-1", "s2", "install share", 14, "%", "unknown", caveat="m", path=led)
    deck = tmp_path / "d.pptx"
    make_deck(deck, "14% of installs, Oct 2025 to Apr 2026. Sources: R-014, R-002 retrieved 2026-10-18")
    assert find_orphans(deck, led) == []
