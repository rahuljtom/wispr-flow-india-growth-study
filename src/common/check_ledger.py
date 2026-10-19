"""Fail if deck/chart numbers have no ledger row (rule 11)."""
import json
import re
import sys
from pathlib import Path

from pptx import Presentation

from src.common.ledger import LEDGER, read_ledger

NUM = re.compile(r"\d[\d,]*\.?\d*")
DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
CLAIM_ID = re.compile(r"\b[A-Z]{1,3}-\d{3}\b")
YEAR = re.compile(r"\b20(?:2[4-7])\b")


def _norm(tok):
    tok = tok.replace(",", "").rstrip(".")
    try:
        return f"{float(tok):g}"
    except ValueError:
        return tok


def _numbers(text):
    text = YEAR.sub(" ", CLAIM_ID.sub(" ", DATE.sub(" ", text)))
    return [_norm(m) for m in NUM.findall(text)]


def _texts(deck):
    for slide in Presentation(str(deck)).slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                yield sh.text_frame.text
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        yield c.text


def find_orphans(deck, ledger_path=LEDGER, charts_dir=None):
    known = set()
    for r in read_ledger(ledger_path):
        if r["value"] != "":
            known.update(_numbers(str(r["value"])))
    found = []
    if deck and Path(deck).exists():
        for t in _texts(deck):
            found += _numbers(t)
    if charts_dir and Path(charts_dir).exists():
        for j in Path(charts_dir).glob("*.json"):
            found += _numbers(json.dumps(json.loads(j.read_text())))
    return sorted({n for n in found if n not in known}, key=found.index)


def main():
    orphans = find_orphans("outputs/deck/deck.pptx", charts_dir="outputs/charts")
    if orphans:
        print("Numbers without ledger rows:", orphans)
        return 1
    print("ledger check ok")
    return 0


if __name__ == "__main__":
    sys.exit(main())
