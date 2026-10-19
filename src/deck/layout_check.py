"""Estimate text clipping and overlaps in outputs/deck/deck.pptx (no renderer available)."""
import sys

from pptx import Presentation
from pptx.util import Emu

from src.deck.layout import block_height_in, table_height_in

SLIDE_H = 7.5
FOOT_TOP = 6.55


def inches(v):
    return Emu(v).inches


def main(path="outputs/deck/deck.pptx"):
    problems = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        for sh in slide.shapes:
            top, left, w, h = inches(sh.top), inches(sh.left), inches(sh.width), inches(sh.height)
            if sh.shape_type == 13:  # picture
                if top + h > FOOT_TOP - 0.05:
                    problems.append((n, "picture overlaps footnote", round(top + h, 2)))
            elif getattr(sh, "has_table", False) and sh.has_table:
                cols = [inches(c.width) for c in sh.table.columns]
                rows = [[c.text for c in r.cells] for r in sh.table.rows]
                pt = min((p.font.size.pt for p in sh.table.cell(1, 0).text_frame.paragraphs if p.font.size), default=11)
                est = table_height_in(rows, cols, pt)
                if top + est > FOOT_TOP - 0.05:
                    problems.append((n, "table may overflow", round(top + est, 2)))
                if sum(cols) > 12.31:
                    problems.append((n, "table wider than slide margin", round(sum(cols), 2)))
            elif sh.has_text_frame and sh.text_frame.text.strip():
                est = 0.0
                for par in sh.text_frame.paragraphs:
                    ppt = next((r.font.size.pt for r in par.runs if r.font.size),
                               par.font.size.pt if par.font.size else 12)
                    gap = 10 if len(sh.text_frame.paragraphs) > 1 and ppt >= 10 else 0
                    est += block_height_in([par.text], w - 0.2, ppt, gap)
                pt = ppt
                limit = h + 0.05
                if est > limit:
                    problems.append((n, "text may clip", f"est {est:.2f}in > box {h:.2f}in at {pt}pt"))
                if top + h > SLIDE_H:
                    problems.append((n, "shape off slide", round(top + h, 2)))
            if left + w > 13.34:
                problems.append((n, "shape beyond right edge", round(left + w, 2)))
    for p in problems:
        print("LAYOUT:", p)
    print("layout check:", "OK" if not problems else f"{len(problems)} issue(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
